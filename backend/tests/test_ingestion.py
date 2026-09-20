from pathlib import Path

from app.ingestion import clean_html, load_bytes, load_text_file, tables_to_markdown


def test_clean_html_removes_script_and_tags() -> None:
    content = "<html><script>alert(1)</script><p>退款处理</p></html>"

    assert "退款处理" in clean_html(content)
    assert "alert" not in clean_html(content)


def test_load_text_file_reads_utf8() -> None:
    file_path = Path("test_ingestion_tmp.txt")
    file_path.write_text("物流应低盐饮食。", encoding="utf-8")

    try:
        text, metadata = load_text_file(file_path)
    finally:
        file_path.unlink(missing_ok=True)

    assert text == "物流应低盐饮食。"
    assert metadata["file_name"] == "test_ingestion_tmp.txt"


def test_load_pdf_returns_text() -> None:
    file_path = Path("test_ingestion_tmp.pdf")
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with file_path.open("wb") as handle:
        writer.write(handle)

    try:
        text, metadata = load_text_file(file_path)
    finally:
        file_path.unlink(missing_ok=True)

    assert isinstance(text, str)
    assert metadata["file_name"] == "test_ingestion_tmp.pdf"


def _pdf_with_text_layer(content: str) -> bytes:
    """构造一个真的带文字层的 PDF（空白页测不出提取链路有没有问题）。"""
    from pypdf import PdfWriter
    from pypdf.generic import (
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
    )

    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 20 150 Td ({content}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)

    from io import BytesIO

    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_load_pdf_extracts_text_layer() -> None:
    # 文本层 PDF 是官方支持路径，必须真的能读出文字（扫描版走 OCR 钩子，不在这一条）
    content = "Refund within 7 days after delivery."

    text, metadata = load_bytes("policy.pdf", _pdf_with_text_layer(content))

    assert content in text
    assert metadata["file_type"] == ".pdf"


def test_tables_to_markdown() -> None:
    markdown = tables_to_markdown(
        [
            [
                ["商品", "退货天数"],
                ["普通商品", "7 天"],
            ]
        ]
    )

    assert "| 商品 | 退货天数 |" in markdown
    assert "| 普通商品 | 7 天 |" in markdown


def test_load_csv_as_markdown_table() -> None:
    # 企业 FAQ 最常见的形态就是 CSV/Excel 表格
    content = "问题,答复\n退货要几天,7 天内可申请\n运费谁承担,质量问题由商家承担\n".encode()

    text, metadata = load_bytes("faq.csv", content)

    assert metadata["file_type"] == ".csv"
    assert "| 问题 | 答复 |" in text
    assert "| 退货要几天 | 7 天内可申请 |" in text


def test_load_xlsx_keeps_sheet_name_and_rows() -> None:
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "售后时效"
    sheet.append(["商品类型", "退货天数"])
    sheet.append(["普通商品", 7])
    sheet.append(["定制商品", "不支持"])
    buffer = BytesIO()
    workbook.save(buffer)

    text, metadata = load_bytes("policy.xlsx", buffer.getvalue())

    assert metadata["file_type"] == ".xlsx"
    assert "## 售后时效" in text
    assert "| 商品类型 | 退货天数 |" in text
    assert "| 定制商品 | 不支持 |" in text


def test_load_docx_keeps_paragraph_and_table_order() -> None:
    from io import BytesIO

    import docx

    document = docx.Document()
    document.add_paragraph("售后政策说明")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "场景"
    table.cell(0, 1).text = "处理方式"
    table.cell(1, 0).text = "定制商品"
    table.cell(1, 1).text = "不支持无理由退货"
    document.add_paragraph("如有疑问请联系人工客服。")
    buffer = BytesIO()
    document.save(buffer)

    text, metadata = load_bytes("policy.docx", buffer.getvalue())

    assert metadata["file_type"] == ".docx"
    assert "售后政策说明" in text
    assert "| 定制商品 | 不支持无理由退货 |" in text
    # 段落在表格前、结尾段落在表格后，说明顺序没乱
    assert text.index("售后政策说明") < text.index("| 定制商品") < text.index("如有疑问")
