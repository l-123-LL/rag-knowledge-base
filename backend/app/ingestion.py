import html
import os
import re
from io import BytesIO
from pathlib import Path

import httpx


def clean_html(content: str) -> str:
    """去掉网页脚本、样式和标签，只保留正文。"""
    content = re.sub(
        r"(?is)<(script|style|noscript)[^>]*>.*?</\1>",
        " ",
        content,
    )
    content = re.sub(r"(?s)<[^>]+>", " ", content)
    return html.unescape(content)


def load_text_file(path: str | Path) -> tuple[str, dict]:
    file_path = Path(path)
    content = file_path.read_bytes()
    return load_bytes(file_path.name, content)


def load_bytes(filename: str, content: bytes) -> tuple[str, dict]:
    suffix = Path(filename).suffix.lower()

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        text = "\n\n".join(
            page.extract_text() or "" for page in reader.pages
        )
        if not text.strip() and os.getenv("OCR_ENABLED", "false").lower() == "true":
            text = ocr_pdf(content)
        table_text = extract_pdf_tables(content)
        if table_text:
            text = f"{text}\n\n{table_text}"
        metadata = {
            "source_path": filename,
            "file_name": filename,
            "file_type": suffix,
        }
        return text, metadata

    if suffix == ".xlsx" or suffix == ".xlsm":
        # 企业资料里最常见的就是 Excel 表格（FAQ 清单、价目、时效表）
        text = xlsx_to_text(content)
        return text, {
            "source_path": filename,
            "file_name": filename,
            "file_type": suffix,
        }

    if suffix == ".docx":
        # Word 里往往是「正文 + 表格」混排，两段都要取
        text = docx_to_text(content)
        return text, {
            "source_path": filename,
            "file_name": filename,
            "file_type": suffix,
        }

    if suffix == ".csv":
        text = csv_to_markdown(content.decode("utf-8", errors="ignore"))
        return text, {
            "source_path": filename,
            "file_name": filename,
            "file_type": suffix,
        }

    text = content.decode("utf-8", errors="ignore")

    if suffix in {".html", ".htm"}:
        text = clean_html(text)

    metadata = {
        "source_path": filename,
        "file_name": filename,
        "file_type": suffix,
    }
    return text, metadata


def csv_to_markdown(text: str) -> str:
    """CSV 转 Markdown 表格：表格结构比裸文本更容易被模型读懂列的含义。"""
    import csv

    rows = list(csv.reader(text.splitlines()))
    return tables_to_markdown([rows]) or text


def xlsx_to_text(content: bytes) -> str:
    """Excel 逐 sheet 转 Markdown 表格，保留 sheet 名作为小节标题。"""
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sections: list[str] = []

    for sheet in workbook.worksheets:
        rows: list[list[str | None]] = []
        for row in sheet.iter_rows(values_only=True):
            values = ["" if cell is None else str(cell).strip() for cell in row]
            if any(values):
                rows.append(values)
        if not rows:
            continue

        table = tables_to_markdown([rows])
        if table:
            sections.append(f"## {sheet.title}\n{table}")

    workbook.close()
    return "\n\n".join(sections)


def docx_to_text(content: bytes) -> str:
    """Word 按文档顺序取出段落与表格（表格转 Markdown）。"""
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(BytesIO(content))
    parts: list[str] = []
    pending_tables: list[list[list[str | None]]] = []

    def flush_tables() -> None:
        if pending_tables:
            parts.append(tables_to_markdown(pending_tables))
            pending_tables.clear()

    # 直接遍历 body 子节点，才能保住「段落 → 表格 → 段落」的原始顺序
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, document).text.strip()
            if text:
                flush_tables()
                parts.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            pending_tables.append(
                [[cell.text.strip() for cell in row.cells] for row in table.rows]
            )

    flush_tables()
    return "\n\n".join(part for part in parts if part.strip())


def tables_to_markdown(tables: list[list[list[str | None]]]) -> str:
    sections: list[str] = []

    for table in tables:
        rows = [
            [str(cell or "").replace("\n", " ").strip() for cell in row]
            for row in table
            if row
        ]
        if len(rows) < 2:
            continue

        header = rows[0]
        body = rows[1:]
        lines = [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join("---" for _ in header) + " |",
        ]
        lines.extend("| " + " | ".join(row) + " |" for row in body)
        sections.append("\n".join(lines))

    return "\n\n".join(sections)


def extract_pdf_tables(content: bytes) -> str:
    """提取文本层 PDF 表格并转为 Markdown，扫描版仍需要 OCR。"""
    try:
        import pdfplumber
    except ImportError:
        return ""

    tables: list[list[list[str | None]]] = []
    with pdfplumber.open(BytesIO(content)) as pdf:
        for page in pdf.pages:
            tables.extend(page.extract_tables() or [])

    return tables_to_markdown(tables)


def ocr_pdf(content: bytes) -> str:
    """可选 OCR：需要安装 pytesseract、pypdfium2 和系统 Tesseract。"""
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError:
        return ""

    pdf = pdfium.PdfDocument(content)
    pages: list[str] = []
    for index in range(len(pdf)):
        page = pdf[index]
        image = page.render(scale=2).to_pil()
        pages.append(pytesseract.image_to_string(image, lang="chi_sim+eng"))
    return "\n\n".join(pages)


def fetch_url_text(url: str, timeout_seconds: float = 20.0) -> tuple[str, dict]:
    """抓取公开网页并清洗正文；尊重服务端错误，但不做 JavaScript 渲染。"""
    response = httpx.get(url, timeout=timeout_seconds, follow_redirects=True)
    response.raise_for_status()
    text = clean_html(response.text)
    metadata = {
        "source_path": url,
        "file_name": url,
        "file_type": ".html",
    }
    return text, metadata
